"""Gera os 4 arquivos do produto (2 versões × [em branco | com exemplo]).

Uso:  python gen/build.py            (gera tudo em dist/)
      python gen/build.py pf-demo    (gera só um)

Processo em 2 passos: (1) gera o arquivo; (2) recalcula uma cópia no LibreOffice, lê os
valores calculados e regera o arquivo final gravando esses valores como cache — assim as
fórmulas aparecem com números mesmo em visualizadores que não recalculam (Prévia do Mac,
celular, e-mail), e o Excel/Sheets recalculam normalmente ao abrir.
"""
import sys, os, shutil, subprocess, json, datetime as dt, glob
sys.path.insert(0, os.path.dirname(__file__))
from lib import *
from ctx import *
import personal, painel, intro, demo

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIST = os.path.join(ROOT, "dist")
import tempfile
SCRATCH = os.environ.get("SCRATCH", os.path.join(tempfile.gettempdir(), "financas_build"))
# recalc.py (LibreOffice) — caminho configurável; sem ele os arquivos saem sem valores em cache (Excel/Sheets recalculam ao abrir)
_rc = os.environ.get("RECALC") or ([*glob.glob("/root/.claude/skills/synced/*/xlsx/scripts/recalc.py")] or [None])[0]
RECALC = _rc

BRAND = {"pf": "Finanças em Dia — Pessoal", "aut": "Finanças em Dia — Autônomo & MEI"}
FILES = {("pf", False): "Financas-em-Dia_Pessoal_EM-BRANCO.xlsx", ("pf", True): "Financas-em-Dia_Pessoal_COM-EXEMPLO.xlsx",
         ("aut", False): "Financas-em-Dia_Autonomo-MEI_EM-BRANCO.xlsx", ("aut", True): "Financas-em-Dia_Autonomo-MEI_COM-EXEMPLO.xlsx"}


def build(kind, demo_on, path, cache=None):
    book = Book(path, cache)
    ctx = Ctx(kind, demo_on, book, BRAND[kind])
    N, L = ctx.N, ctx.L
    book.wb.set_properties({"title": BRAND[kind], "subject": "Planilha de controle financeiro",
                            "author": "Finanças em Dia", "comments": "Compatível com Excel e Google Sheets"})
    personal.config(ctx)
    personal.dados(ctx)
    rows = demo.personal_rows(aut=ctx.aut) if demo_on else None
    metas_names = X(N["METAS"], RNG(L["m_r0"], 2, L["m_r0"] + NGOAL - 1, 2))
    personal.lancamentos(
        ctx, "LANC", area="pessoal",
        title="📝 Lançamentos — registre cada entrada e saída" if not ctx.aut else "📝 Lançamentos pessoais",
        subtitle="Preencha as colunas amarelas. Valor sempre positivo; o Tipo (Entrada/Saída) é automático pela categoria.",
        ent_key="c_ent", sai_key="c_sai", all_col_key="d_all_c", extra_label="Meta (opcional)",
        extra_list_ref=metas_names, demo_rows=rows)
    personal.resumo(ctx)
    personal.orcamento(ctx)
    personal.metas(ctx)
    painel.painel_pessoal(ctx)
    if ctx.aut:
        import business
        business.build_all(ctx)
    intro.inicio(ctx)
    ctx.ws["INICIO"].activate()
    ctx.ws["INICIO"].set_first_sheet()
    book.close()
    return ctx


def serial(v):
    if isinstance(v, dt.datetime):
        d = v - dt.datetime(1899, 12, 30)
        return d.days + d.seconds / 86400
    if isinstance(v, dt.date):
        return (v - dt.date(1899, 12, 30)).days
    return v


def compute_cache(kind, demo_on, tag):
    if not RECALC:
        print("AVISO: recalc.py não encontrado — gerando sem valores em cache.")
        return {}, {}
    """Passo 1: gera, recalcula uma cópia no LO e devolve {(aba, 'A1'): valor} + erros."""
    from openpyxl import load_workbook
    os.makedirs(SCRATCH, exist_ok=True)
    p1 = os.path.join(SCRATCH, f"{tag}_pass1.xlsx")
    build(kind, demo_on, p1)
    p2 = os.path.join(SCRATCH, f"{tag}_recalc.xlsx")
    shutil.copy(p1, p2)
    out = subprocess.run([sys.executable, RECALC, p2, "180"], capture_output=True, text=True,
                         cwd=os.path.dirname(RECALC))
    try:
        rep = json.loads(out.stdout)
    except Exception:
        raise SystemExit("recalc falhou: " + out.stdout + out.stderr)
    if "error" in rep:
        raise SystemExit("recalc falhou: " + json.dumps(rep))
    wf = load_workbook(p1)                  # fórmulas
    wv = load_workbook(p2, data_only=True)  # valores
    cache = {}
    for ws in wf.worksheets:
        wsv = wv[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    cache[(ws.title, c.coordinate)] = serial(wsv[c.coordinate].value)
    return cache, rep


def make(kind, demo_on):
    tag = f"{kind}_{'demo' if demo_on else 'blank'}"
    cache, rep = compute_cache(kind, demo_on, tag)
    os.makedirs(DIST, exist_ok=True)
    final = os.path.join(DIST, FILES[(kind, demo_on)])
    build(kind, demo_on, final, cache)
    print(tag, "→", final, "| fórmulas:", rep.get("total_formulas"), "| erros:", rep.get("total_errors"))
    if rep.get("total_errors"):
        print(json.dumps(rep.get("error_summary"), ensure_ascii=False, indent=1)[:3000])
    return rep


if __name__ == "__main__":
    targets = sys.argv[1:] or ["pf-blank", "pf-demo", "aut-blank", "aut-demo"]
    for t in targets:
        k, d = t.split("-")
        make(k, d == "demo")
