"""Testes de interatividade: altera entradas nos arquivos de exemplo, recalcula no LibreOffice e confere."""
import sys, os, shutil, subprocess, json, glob, datetime as dt
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from openpyxl import load_workbook
from lib import *
from ctx import *
from build import FILES, SCRATCH, RECALC, DIST
import demo

fails = []

def chk(name, got, exp, tol=0.011):
    ok = got is not None and (abs(float(got) - float(exp)) <= tol if not isinstance(exp, str) else got == exp)
    if not ok: fails.append(name)
    print(("  ok    " if ok else "  FALHA "), name, "| obtido:", got, "| esperado:", exp)

def run(kind, edits, tag):
    src = os.path.join(DIST, FILES[(kind, True)])
    wb = load_workbook(src)
    edits(wb)
    p = os.path.join(SCRATCH, f"sc_{tag}.xlsx")
    wb.save(p)
    out = subprocess.run([sys.executable, RECALC, p, "180"], capture_output=True, text=True, cwd=os.path.dirname(RECALC))
    rep = json.loads(out.stdout)
    print(f"[{tag}] fórmulas={rep.get('total_formulas')} erros={rep.get('total_errors')}", rep.get("error_summary") or "")
    if rep.get("total_errors"): fails.append(tag + ":erros")
    return load_workbook(p, data_only=True)

def ctx_for(kind):
    return Ctx(kind, True, Book("/tmp/_s.xlsx"), "x")

# ---- PF
c = ctx_for("pf"); L, N = c.L, c.N
rows = demo.personal_rows(aut=False)
def edit_pf(wb):
    wb[N["PAINEL"]].cell(row=L["p_sel"][0], column=4).value = "Julho"
    ws = wb[N["LANC"]]
    r = L["l_r0"] + len(rows)               # primeira linha livre
    ws.cell(row=r, column=2).value = dt.datetime(2026, 7, 31)
    ws.cell(row=r, column=3).value = "Teste"
    ws.cell(row=r, column=4).value = "Lazer"
    ws.cell(row=r, column=6).value = 1000
    wb[N["CONFIG"]].cell(row=L["cfg_meta"][0], column=3).value = 0.30
wbv = run("pf", edit_pf, "pf_jul")
E7 = sum(v for d,_,cat,_,v,_,_ in rows if d.month==7 and cat in __import__("personal").ENT_PF)
S7 = sum(v for d,_,cat,_,v,_,_ in rows if d.month==7 and cat not in __import__("personal").ENT_PF) + 1000
ws = wbv[N["PAINEL"]]
chk("PF julho: entradas", ws.cell(row=L["p_sel"][0]+3, column=2).value, E7)
chk("PF julho: saídas (inclui novo lançamento de 1000)", ws.cell(row=L["p_sel"][0]+3, column=4).value, S7)
chk("PF julho: saldo negativo", ws.cell(row=L["p_sel"][0]+3, column=6).value, E7 - S7)
print("    alerta 1:", ws.cell(row=L["p_sel"][0]+ 3 + 3 + 16*2 + 2 + 9 + 1, column=2).value)
chk("PF: categoria renomeada no Resumo (Config)", wbv[N["RESUMO"]].cell(row=L["ra_sai0"]+6, column=2).value, "Lazer")

def edit_rename(wb):
    wb[N["CONFIG"]].cell(row=L["ls"]+1+6, column=L["c_sai"]).value = "Diversão"
wbv = run("pf", edit_rename, "pf_rename")
ws = wbv[N["LANC"]]
tipos = [ws.cell(row=r, column=7).value for r in range(L["l_r0"], L["l_r0"] + len(rows))]
chk("PF renomear categoria: lançamentos antigos viram 'Verificar'", tipos.count("Verificar"), sum(1 for r in rows if r[2]=="Lazer"))
print("    verificação:", ws.cell(row=c.CS+1, column=7).value)
chk("PF renomear: Resumo mostra nome novo", wbv[N["RESUMO"]].cell(row=L["ra_sai0"]+6, column=2).value, "Diversão")

def edit_year(wb):
    wb[N["CONFIG"]].cell(row=L["cfg_ano"][0], column=3).value = 2027
wbv = run("pf", edit_year, "pf_2027")
chk("PF ano 2027: entradas de set = 0", wbv[N["PAINEL"]].cell(row=L["p_sel"][0]+3, column=2).value, 0)

# ---- AUT
c = ctx_for("aut"); L, N = c.L, c.N
rec, lan = demo.business_rows()
def edit_aut(wb):
    wb[N["IMP"]].cell(row=L["i_p0"], column=5).value = "Simples Nacional"
    ws = wb[N["REC"]]
    # paga a cobrança do Studio Aurora de setembro (a 1ª pendente) em 30/09
    for r in range(L["r_r0"], L["r_r0"] + len(rec)):
        if ws.cell(row=r, column=3).value == "Studio Aurora" and ws.cell(row=r, column=7).value is None:
            ws.cell(row=r, column=7).value = dt.datetime(2026, 9, 30); break
    wb[N["PNEG"]].cell(row=L["pn_sel"][0], column=4).value = "Agosto"
wbv = run("aut", edit_aut, "aut_simples")
ws = wbv[N["IMP"]]
chk("AUT Simples: imposto jan = 6% da receita jan", ws.cell(row=L["i_r0"], column=4).value, round(3300*0.06, 2))
print("    msg limite:", wbv[N["IMP"]].cell(row=c.CS+6, column=8).value)
chk("AUT: recebido no ano sobe após pagar (55530+1800)", wbv[N["REC"]].cell(row=c.CS+1, column=5).value, 55530 + 1800)
chk("AUT: em atraso cai (4430-1800)", wbv[N["REC"]].cell(row=c.CS+1, column=4).value, 4430 - 1800)
chk("AUT: Painel Negócio mês de agosto: receita", wbv[N["PNEG"]].cell(row=c.CS+3, column=2).value, 7300)

print("\nRESULTADO:", "TUDO CONFERE" if not fails else f"{len(fails)} FALHA(S): {fails}")
sys.exit(1 if fails else 0)
