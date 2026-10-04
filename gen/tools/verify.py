"""Confere os valores calculados nos arquivos de exemplo contra um cálculo independente em Python.
Uso: python gen/tools/verify.py   (lê dist/*.xlsx — valores em cache gravados pelo build)"""
import sys, os, collections, datetime as dt
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from openpyxl import load_workbook
from lib import *
from ctx import *
import demo
from personal import ENT_PF, ENT_AUT, SAI, NSAI_L, NENT_L
from build import FILES

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DIST = os.path.join(ROOT, "dist")
fails = []


def chk(name, got, exp, tol=0.011):
    ok = (got is not None) and abs(float(got) - float(exp)) <= tol
    if not ok:
        fails.append(name)
    print(("  ok   " if ok else "  FALHA"), name, "| planilha:", got, "| esperado:", round(float(exp), 2))


def ctx_for(kind):
    b = Book("/tmp/_v.xlsx")
    c = Ctx(kind, True, b, "x")
    return c


def cell(wb, sheet, r, c):
    return wb[sheet].cell(row=r, column=c).value


def verify_personal(kind):
    aut = kind == "aut"
    ctx = ctx_for(kind)
    L, N = ctx.L, ctx.N
    wb = load_workbook(os.path.join(DIST, FILES[(kind, True)]), data_only=True)
    rows = demo.personal_rows(aut=aut)
    ent_names = set(ENT_AUT if aut else ENT_PF)
    E = collections.Counter(); S = collections.Counter(); CAT = collections.Counter(); SAV = collections.Counter()
    meta = collections.Counter()
    for d, desc, cat, f, v, m, o in rows:
        if cat in ent_names:
            E[d.month] += v
        else:
            S[d.month] += v
            CAT[(cat, d.month)] += v
            if cat == "Reserva e Metas":
                SAV[d.month] += v
        if m:
            meta[m] += v
    print(f"[{kind}] Resumo anual (pessoal)")
    for m in (1, 4, 7, 9):
        chk(f"entradas mês {m}", cell(wb, N["RESUMO"], L["ra_enttot"], 2 + m), E[m])
        chk(f"saídas mês {m}", cell(wb, N["RESUMO"], L["ra_saitot"], 2 + m), S[m])
    saldo0 = 2500
    acum9 = saldo0 + sum(E[m] - S[m] for m in range(1, 10))
    chk("saldo acumulado em set", cell(wb, N["RESUMO"], L["ra_acum"], 11), acum9)
    chk("taxa de poupança set", cell(wb, N["RESUMO"], L["ra_taxa"], 11), (E[9] - S[9] + SAV[9]) / E[9], 0.0005)
    i_lazer = SAI.index("Lazer")
    chk("Lazer em jul", cell(wb, N["RESUMO"], L["ra_sai0"] + i_lazer, 9), CAT[("Lazer", 7)])
    print(f"[{kind}] Painel (set)")
    chk("KPI entradas", cell(wb, N["PAINEL"], L["p_sel"][0] + 3, 2), E[9])
    chk("KPI saídas", cell(wb, N["PAINEL"], L["p_sel"][0] + 3, 4), S[9])
    chk("KPI saldo", cell(wb, N["PAINEL"], L["p_sel"][0] + 3, 6), E[9] - S[9])
    chk("KPI saldo acumulado", cell(wb, N["PAINEL"], L["p_sel"][0] + 3, 10), acum9)
    top = max((CAT[(c, 9)], c) for c in SAI)
    chk("maior gasto (valor)", cell(wb, N["PAINEL"], L["p_sel"][0] + 4, 12), top[0])
    print("    maior gasto:", cell(wb, N["PAINEL"], L["p_sel"][0] + 3, 12), "| esperado:", top[1])
    print(f"[{kind}] Orçamento (set) e Metas")
    chk("orçamento: realizado Alimentação", cell(wb, N["ORC"], L["o_c0"] + SAI.index("Alimentação"), 4), CAT[("Alimentação", 9)])
    chk("orçamento: total realizado", cell(wb, N["ORC"], L["o_tot"], 4), S[9])
    gl = [g for g in (__import__("personal").GOALS_DEMO_AUT if aut else __import__("personal").GOALS_DEMO)]
    for i, g in enumerate(gl):
        chk(f"meta '{g[0]}' total guardado", cell(wb, N["METAS"], L["m_r0"] + i, 7), g[2] + meta[g[0]])
    # tipos
    ws = wb[N["LANC"]]
    tipos = collections.Counter(ws.cell(row=r, column=7).value for r in range(L["l_r0"], L["l_r0"] + len(rows)))
    n_ent = sum(1 for r in rows if r[2] in ent_names)
    chk("nº de lançamentos tipo Entrada", tipos["Entrada"], n_ent)
    chk("nº de lançamentos tipo Saída", tipos["Saída"], len(rows) - n_ent)
    chk("nenhum 'Verificar'", tipos.get("Verificar", 0), 0)
    return wb, ctx, E, S


def verify_business():
    ctx = ctx_for("aut")
    L, N = ctx.L, ctx.N
    wb = load_workbook(os.path.join(DIST, FILES[("aut", True)]), data_only=True)
    rec, lan = demo.business_rows()
    ref = dt.date(2026, 9, 30)
    print("[aut] Recebimentos / Clientes")
    st = collections.Counter(wb[N["REC"]].cell(row=r, column=9).value for r in range(L["r_r0"], L["r_r0"] + len(rec)))
    exp_late = [r for r in rec if r[5] is None and r[4] < ref]
    exp_pend = [r for r in rec if r[5] is None and r[4] >= ref]
    chk("recebimentos 'Atrasado'", st["Atrasado"], len(exp_late))
    chk("recebimentos 'Pendente'", st["Pendente"], len(exp_pend))
    chk("recebimentos 'Pago'", st["Pago"], sum(1 for r in rec if r[5]))
    chk("KPI em aberto", cell(wb, N["REC"], ctx.CS + 1, 2), sum(r[3] for r in rec if r[5] is None))
    chk("KPI em atraso", cell(wb, N["REC"], ctx.CS + 1, 4), sum(r[3] for r in exp_late))
    paid = sum(r[3] for r in rec if r[5])
    chk("KPI recebido no ano", cell(wb, N["REC"], ctx.CS + 1, 5), paid)
    chk("clientes: total recebido", cell(wb, N["CLI"], L["cli_tot"], 8), paid)
    chk("clientes: total a receber", cell(wb, N["CLI"], L["cli_tot"], 9), sum(r[3] for r in rec if r[5] is None))
    aur = sum(r[3] for r in rec if r[1] == "Studio Aurora" and r[5])
    chk("cliente Studio Aurora recebido", cell(wb, N["CLI"], L["cli_r0"], 8), aur)

    print("[aut] Impostos")
    pay_m = collections.Counter(); avu_m = collections.Counter(); opex = collections.Counter(); pro = collections.Counter()
    for e, c, d, v, venc, pg, f in rec:
        if pg:
            pay_m[pg.month] += v
    for d, desc, cat, f, v, x, o in lan:
        if cat == NENT_L[0]:
            avu_m[d.month] += v
        elif cat == NSAI_L[0]:
            pro[d.month] += v
        elif cat in NSAI_L:
            opex[d.month] += v
    for m in (1, 6, 9):
        chk(f"IMP receita recebida mês {m}", cell(wb, N["IMP"], L["i_r0"] + m - 1, 3), pay_m[m] + avu_m[m])
    das = 1621 * 0.05 + 5
    chk("DAS mensal", cell(wb, N["IMP"], L["i_r0"], 4), das)
    chk("IMP status set = 'A pagar'", 1 if cell(wb, N["IMP"], L["i_r0"] + 8, 9) == "A pagar" else 0, 1)
    chk("IMP status ago = 'Pago'", 1 if cell(wb, N["IMP"], L["i_r0"] + 7, 9) == "Pago" else 0, 1)
    acum = sum(pay_m[m] + avu_m[m] for m in range(1, 10))
    chk("receita acumulada (limite MEI)", cell(wb, N["IMP"], L["i_r0"] + 8, 12), acum)
    chk("% do limite", cell(wb, N["IMP"], ctx.CS + 3, 12), acum / 81000, 0.0005)
    chk("projeção anual", cell(wb, N["IMP"], ctx.CS + 4, 12), acum / 9 * 12)

    print("[aut] Fluxo de caixa")
    R_ = L["f_rows"] if "f_rows" in L else None
    f0 = L["f_0"]
    rows_ = dict(ini=f0, rec=f0 + 1, avu=f0 + 2, out=f0 + 3, nov=f0 + 4, ent=f0 + 5, dsp=f0 + 6, imp=f0 + 7, pro=f0 + 8,
                 sai=f0 + 9, var=f0 + 10, fim=f0 + 11, res=f0 + 12, mar=f0 + 13)
    saldo = 4000
    imp_paid = collections.Counter()
    for i in range(8):
        imp_paid[i + 2] += das       # pago de fev a set (competências jan-ago)
    for m in range(1, 10):
        saldo += pay_m[m] + avu_m[m] - opex[m] - imp_paid[m] - pro[m]
        if m in (1, 5, 9):
            chk(f"saldo final mês {m}", cell(wb, N["FLUXO"], rows_["fim"], 2 + m), saldo)
    avg_rev = sum(pay_m[m] + avu_m[m] for m in range(1, 10)) / 9
    avg_opex = sum(opex[m] for m in range(1, 10)) / 9
    chk("receita média realizada", cell(wb, N["FLUXO"], L["f_prem"] + 4, 3), avg_rev)
    chk("despesa média", cell(wb, N["FLUXO"], L["f_prem"] + 2, 3), avg_opex)
    open_ = sum(r[3] for r in rec if r[5] is None)
    chk("out: recebimentos previstos (rola em aberto)", cell(wb, N["FLUXO"], rows_["rec"], 12), open_)
    chk("out: receita nova prevista", cell(wb, N["FLUXO"], rows_["nov"], 12), avg_rev * 0.8)
    chk("out: despesas (média)", cell(wb, N["FLUXO"], rows_["dsp"], 12), avg_opex)
    chk("out: DAS em aberto (set)", cell(wb, N["FLUXO"], rows_["imp"], 12), das)
    chk("out: pró-labore previsto", cell(wb, N["FLUXO"], rows_["pro"], 12), 4500)
    s_out = saldo + open_ + avg_rev * 0.8 - avg_opex - das - 4500
    chk("saldo final out (projetado)", cell(wb, N["FLUXO"], rows_["fim"], 12), s_out)
    s_nov = s_out + avg_rev * 0.8 - avg_opex - das - 4500
    chk("saldo final nov (projetado)", cell(wb, N["FLUXO"], rows_["fim"], 13), s_nov)
    print("[aut] Painel Negócio / Geral (set)")
    ser = L["b_ser"]
    chk("Dados: receita set", cell(wb, N["DADOS"], ser["rec"], 11), pay_m[9] + avu_m[9])
    chk("Dados: lucro set", cell(wb, N["DADOS"], ser["lucro"], 11), pay_m[9] + avu_m[9] - opex[9] - imp_paid[9])
    chk("KPI neg: receita do mês", cell(wb, N["PNEG"], ctx.CS + 3, 2), pay_m[9] + avu_m[9])
    chk("KPI neg: a receber", cell(wb, N["PNEG"], ctx.CS + 3, 8), open_)
    chk("KPI geral: caixa negócio", cell(wb, N["GERAL"], ctx.CS + 3, 4), saldo)


for kind in ("pf", "aut"):
    path = os.path.join(DIST, FILES[(kind, True)])
    if not os.path.exists(path):
        continue
    sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    verify_personal(kind)
if os.path.exists(os.path.join(DIST, FILES[("aut", True)])):
    verify_business()
print("\nRESULTADO:", "TUDO CONFERE" if not fails else f"{len(fails)} FALHA(S): {fails}")
sys.exit(1 if fails else 0)
